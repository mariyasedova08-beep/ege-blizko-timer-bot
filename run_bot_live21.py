import io
import json
import re
import sqlite3
from datetime import datetime

from openpyxl import load_workbook

import run_bot_live20

live20 = run_bot_live20
bot = live20.bot

_original_import_students_document = bot.import_students_document


def _text(value):
    return str(value or "").strip()


def _is_monitoring_export(rows):
    if not rows:
        return False
    header = [str(v or "").strip().lower() for v in rows[0]]
    return (
        any("пройдено уроков" in x for x in header)
        and any("id ученика" in x for x in header)
        and any("результат" in x for x in header)
    )


def _lesson_date_iso(lesson_name):
    # В мониторинге CoreApp нет отдельной даты завершения каждого урока,
    # поэтому для старых уроков берём дату, указанную в названии урока.
    match = re.search(r"(?<!\d)(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?", lesson_name or "")
    if not match:
        return None
    day = int(match.group(1))
    month = int(match.group(2))
    year_raw = match.group(3)
    if year_raw:
        year = int(year_raw)
        if year < 100:
            year += 2000
    else:
        # Годовой курс 2026/27: сентябрь-декабрь относятся к 2026,
        # январь-июнь — к 2027.
        year = 2026 if month >= 7 else 2027
    try:
        return datetime(year, month, day, 12, 0, tzinfo=bot.TIMEZONE).isoformat()
    except ValueError:
        return None


def _lesson_key(lesson_name):
    m = re.search(r"урок\s*№\s*(\d+)", (lesson_name or "").lower())
    if m:
        return f"monitoring:lesson:{m.group(1)}"
    normalized = re.sub(r"[^a-zа-я0-9]+", "-", (lesson_name or "").lower().replace("ё", "е")).strip("-")
    return "monitoring:" + normalized[:120]


def _score_counts(result_value, raw_correct, raw_total):
    # Экспорт CoreApp содержит готовый процент результата. Он точнее для
    # исторической статистики, чем соседние числовые колонки (например,
    # в выгрузке встречается 14/14 при результате 99%).
    result_text = _text(result_value).replace(",", ".")
    m = re.search(r"(-?\d+(?:\.\d+)?)\s*%", result_text)
    if m:
        return m.group(1), "100"

    def num(v):
        s = _text(v).replace(",", ".")
        try:
            return str(float(s)).rstrip("0").rstrip(".")
        except ValueError:
            return ""

    return num(raw_correct), num(raw_total)


def _student_name_from_db(conn, user_id, email, fallback):
    row = None
    if user_id:
        row = conn.execute(
            """
            SELECT coalesce(nullif(display_name, ''), nullif(user_name, ''), user_email, '')
            FROM students WHERE coreapp_user_id = ? LIMIT 1
            """,
            (user_id,),
        ).fetchone()
    if row is None and email:
        row = conn.execute(
            """
            SELECT coalesce(nullif(display_name, ''), nullif(user_name, ''), user_email, '')
            FROM students WHERE lower(user_email) = lower(?) LIMIT 1
            """,
            (email,),
        ).fetchone()
    return (row[0] if row and row[0] else fallback) or "Ученик"


def import_monitoring_results(file_bytes):
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not _is_monitoring_export(rows):
        return None

    imported = 0
    duplicates = 0
    skipped_undated = 0
    completed_found = 0
    students_with_completed = set()
    current = None

    with sqlite3.connect(bot.COREAPP_DB_PATH) as conn:
        for row in rows[1:]:
            values = list(row) + [None] * 7
            a, b, c, d, e, f, g = values[:7]
            col_a, col_b, col_c, status = _text(a), _text(b), _text(c), _text(d).lower()

            # Строка ученика: во второй колонке e-mail, в третьей — ID CoreApp.
            if "@" in col_b and col_c:
                current = {
                    "name": col_a,
                    "email": bot.normalize_email(col_b),
                    "user_id": col_c,
                }
                continue

            if not current or status != "complete" or not col_b:
                continue

            completed_found += 1
            received_at = _lesson_date_iso(col_b)
            if not received_at:
                skipped_undated += 1
                continue

            lesson_id = _lesson_key(col_b)
            email = current["email"]
            user_id = current["user_id"]
            shown_name = _student_name_from_db(conn, user_id, email, current["name"])
            correct_count, total_count = _score_counts(f, a, c)

            # Обновляем только технический ID CoreApp, не трогаем отображаемое имя.
            existing_student = conn.execute(
                "SELECT id FROM students WHERE coreapp_user_id = ? OR lower(user_email) = lower(?) LIMIT 1",
                (user_id, email),
            ).fetchone()
            if existing_student:
                conn.execute(
                    """
                    UPDATE students
                    SET coreapp_user_id = CASE WHEN coalesce(coreapp_user_id, '') = '' THEN ? ELSE coreapp_user_id END,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (user_id, datetime.now(bot.TIMEZONE).isoformat(), existing_student[0]),
                )

            duplicate = conn.execute(
                """
                SELECT 1 FROM homework_submissions
                WHERE (user_id = ? OR lower(user_email) = lower(?))
                  AND lesson_name = ?
                LIMIT 1
                """,
                (user_id, email, col_b),
            ).fetchone()
            if duplicate:
                duplicates += 1
                continue

            payload = {
                "source": "coreapp_monitoring_xlsx",
                "user_id": user_id,
                "user_email": email,
                "user_name": shown_name,
                "lesson_id": lesson_id,
                "lesson_name": col_b,
                "correct_count": correct_count,
                "total_count": total_count,
                "result": _text(f),
            }
            conn.execute(
                """
                INSERT INTO homework_submissions (
                    received_at, user_id, user_email, user_name, course_id,
                    lesson_id, lesson_name, correct_count, total_count, raw_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    received_at,
                    user_id,
                    email,
                    shown_name,
                    "",
                    lesson_id,
                    col_b,
                    correct_count,
                    total_count,
                    json.dumps(payload, ensure_ascii=False),
                ),
            )
            imported += 1
            students_with_completed.add(email or user_id)

        conn.commit()

    return {
        "completed_found": completed_found,
        "imported": imported,
        "duplicates": duplicates,
        "skipped_undated": skipped_undated,
        "students": len(students_with_completed),
    }


async def smart_xlsx_import(update, context):
    if update.effective_chat.type != "private" or not bot.user_is_admin(update):
        return
    document = update.message.document
    if not document or not (document.file_name or "").lower().endswith(".xlsx"):
        return await _original_import_students_document(update, context)

    telegram_file = await document.get_file()
    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)
    file_bytes = buffer.getvalue()

    try:
        result = import_monitoring_results(file_bytes)
    except Exception as exc:
        print("Monitoring XLSX import error:", repr(exc))
        await update.message.reply_text("Не получилось прочитать выгрузку мониторинга CoreApp.")
        return

    if result is None:
        return await _original_import_students_document(update, context)

    lines = [
        "✅ Старые завершённые уроки из CoreApp загружены.",
        f"Добавлено сдач: {result['imported']}",
        f"Учеников со сдачами: {result['students']}",
    ]
    if result["duplicates"]:
        lines.append(f"Уже были в базе: {result['duplicates']}")
    if result["skipped_undated"]:
        lines.append(f"Без даты в названии урока пропущено: {result['skipped_undated']}")
    lines.extend([
        "",
        "Теперь проверь: /monthstats 09.2026",
    ])
    await update.message.reply_text("\n".join(lines))


# В существующем main обработчик документов берёт функцию из bot-модуля,
# поэтому подменяем только импорт XLSX и сохраняем всю остальную логику бота.
bot.import_students_document = smart_xlsx_import


if __name__ == "__main__":
    live20.main()
