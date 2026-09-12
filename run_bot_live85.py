import html
import io
import sqlite3
from datetime import datetime

from openpyxl import load_workbook
from telegram import InlineKeyboardButton, InlineKeyboardMarkup

import run_bot_live84

live84 = run_bot_live84
live79 = live84.live79
bot = live79.bot
live23 = live79.live23


PAYMENT_SHEET_NAME = "11 КЛАСС"
PAYMENT_PERIODS = (
    ("2026-09", "Сентябрь"),
    ("2026-10", "Октябрь"),
    ("2026-11", "Ноябрь"),
    ("2026-12", "Декабрь"),
    ("2027-01", "Январь"),
    ("2027-02", "Февраль"),
    ("2027-03", "Март"),
    ("2027-04", "Апрель"),
    ("2027-05", "Май"),
)
PAYMENT_PERIOD_BY_MONTH = {name.upper(): period for period, name in PAYMENT_PERIODS}


def ensure_payment_tables():
    with sqlite3.connect(bot.COREAPP_DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS payment_students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_key TEXT NOT NULL UNIQUE,
                display_name TEXT NOT NULL,
                monthly_amount REAL,
                source_sheet TEXT NOT NULL DEFAULT '11 КЛАСС',
                active INTEGER NOT NULL DEFAULT 1,
                student_messages_enabled INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS payment_coverage (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                payment_student_id INTEGER NOT NULL,
                period TEXT NOT NULL,
                amount REAL NOT NULL,
                imported_at TEXT NOT NULL,
                UNIQUE(payment_student_id, period),
                FOREIGN KEY(payment_student_id) REFERENCES payment_students(id)
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS payment_import_status (
                id INTEGER PRIMARY KEY CHECK(id = 1),
                imported_at TEXT,
                source_filename TEXT,
                student_count INTEGER NOT NULL DEFAULT 0,
                coverage_count INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute(
            """
            INSERT OR IGNORE INTO payment_import_status
                (id, student_count, coverage_count)
            VALUES (1, 0, 0)
            """
        )
        conn.commit()


def _clean_name(value):
    return " ".join(str(value or "").split()).strip()


def _name_key(value):
    return _clean_name(value).casefold()


def _amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        cleaned = str(value).replace("₽", "").replace(" ", "").replace(",", ".")
        try:
            number = float(cleaned)
        except ValueError:
            return None
    return number if number > 0 else None


def _payment_sheet(workbook):
    for name in workbook.sheetnames:
        if str(name).strip().casefold() == PAYMENT_SHEET_NAME.casefold():
            return workbook[name]
    return None


def parse_payment_workbook(file_bytes):
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = _payment_sheet(workbook)
    if sheet is None:
        return None

    # Некоторые выгрузки Google Sheets не содержат корректный XML-диапазон,
    # поэтому max_row/max_column могут быть None. Читаем только небольшой
    # безопасный прямоугольник и определяем границы по заголовкам.
    matrix = [
        list(row)
        for row in sheet.iter_rows(min_row=1, max_row=50, max_col=100, values_only=True)
    ]
    if not matrix:
        raise ValueError("вкладка «11 КЛАСС» пуста")

    first_header = _clean_name(matrix[0][0]).upper()
    if first_header != "МЕСЯЦ":
        raise ValueError("на вкладке «11 КЛАСС» в A1 ожидается заголовок «МЕСЯЦ»")

    total_column = None
    for column, value in enumerate(matrix[0][1:], 1):
        if _clean_name(value).upper() == "ИТОГ":
            total_column = column
            break
    if total_column is None:
        raise ValueError("на вкладке «11 КЛАСС» не найден столбец «ИТОГ»")

    rows_by_period = {}
    for row in matrix[1:]:
        month_name = _clean_name(row[0]).upper()
        period = PAYMENT_PERIOD_BY_MONTH.get(month_name)
        if period:
            rows_by_period[period] = row
    missing_months = [name for period, name in PAYMENT_PERIODS if period not in rows_by_period]
    if missing_months:
        raise ValueError("не найдены месяцы: " + ", ".join(missing_months))

    students = []
    seen_keys = set()
    for column in range(1, total_column):
        name = _clean_name(matrix[0][column])
        if not name:
            continue
        key = _name_key(name)
        if key in seen_keys:
            raise ValueError(f"имя «{name}» встречается дважды")
        seen_keys.add(key)
        coverage = {}
        for period, _month_name in PAYMENT_PERIODS:
            value = _amount(rows_by_period[period][column])
            if value is not None:
                coverage[period] = value
        if coverage:
            monthly_amount = coverage.get(PAYMENT_PERIODS[0][0]) or next(iter(coverage.values()))
            students.append(
                {
                    "source_key": key,
                    "display_name": name,
                    "monthly_amount": monthly_amount,
                    "coverage": coverage,
                }
            )
    if not students:
        raise ValueError("не найдено ни одной оплаты")
    return students


def save_payment_snapshot(students, source_filename):
    ensure_payment_tables()
    now = datetime.now(bot.TIMEZONE).isoformat()
    active_keys = [student["source_key"] for student in students]
    coverage_count = 0
    with sqlite3.connect(bot.COREAPP_DB_PATH) as conn:
        for student in students:
            conn.execute(
                """
                INSERT INTO payment_students
                    (source_key, display_name, monthly_amount, source_sheet, active,
                     student_messages_enabled, updated_at)
                VALUES (?, ?, ?, ?, 1, 0, ?)
                ON CONFLICT(source_key) DO UPDATE SET
                    display_name = excluded.display_name,
                    monthly_amount = excluded.monthly_amount,
                    source_sheet = excluded.source_sheet,
                    active = 1,
                    updated_at = excluded.updated_at
                """,
                (
                    student["source_key"],
                    student["display_name"],
                    student["monthly_amount"],
                    PAYMENT_SHEET_NAME,
                    now,
                ),
            )
            student_id = conn.execute(
                "SELECT id FROM payment_students WHERE source_key = ?",
                (student["source_key"],),
            ).fetchone()[0]
            conn.execute(
                "DELETE FROM payment_coverage WHERE payment_student_id = ? AND period BETWEEN '2026-09' AND '2027-05'",
                (student_id,),
            )
            for period, amount in student["coverage"].items():
                conn.execute(
                    """
                    INSERT INTO payment_coverage
                        (payment_student_id, period, amount, imported_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (student_id, period, amount, now),
                )
                coverage_count += 1

        if active_keys:
            placeholders = ",".join("?" for _ in active_keys)
            conn.execute(
                f"""
                UPDATE payment_students
                SET active = 0, updated_at = ?
                WHERE source_sheet = ? AND source_key NOT IN ({placeholders})
                """,
                (now, PAYMENT_SHEET_NAME, *active_keys),
            )
        conn.execute(
            """
            UPDATE payment_import_status
            SET imported_at = ?, source_filename = ?, student_count = ?, coverage_count = ?
            WHERE id = 1
            """,
            (now, _clean_name(source_filename)[:200], len(students), coverage_count),
        )
        conn.commit()
    return len(students), coverage_count


def _current_course_period():
    today = bot.today_moscow()
    period = f"{today.year:04d}-{today.month:02d}"
    return period if period in dict(PAYMENT_PERIODS) else None


def _completed_course_periods():
    today = bot.today_moscow()
    first_year, first_month = map(int, PAYMENT_PERIODS[0][0].split("-"))
    last_year, last_month = map(int, PAYMENT_PERIODS[-1][0].split("-"))
    if (today.year, today.month) < (first_year, first_month):
        return []
    if (today.year, today.month) > (last_year, last_month):
        return [period for period, _name in PAYMENT_PERIODS]
    current = f"{today.year:04d}-{today.month:02d}"
    return [period for period, _name in PAYMENT_PERIODS if period < current]


def _money(value):
    number = float(value or 0)
    if number.is_integer():
        return f"{int(number):,}".replace(",", " ") + " ₽"
    return f"{number:,.2f}".replace(",", " ").replace(".", ",") + " ₽"


def _payment_rows():
    ensure_payment_tables()
    with sqlite3.connect(bot.COREAPP_DB_PATH) as conn:
        students = conn.execute(
            """
            SELECT id, display_name, monthly_amount, student_messages_enabled
            FROM payment_students
            WHERE active = 1
            ORDER BY lower(display_name)
            """
        ).fetchall()
        result = []
        for student_id, name, monthly_amount, messages_enabled in students:
            coverage = dict(
                conn.execute(
                    """
                    SELECT period, amount
                    FROM payment_coverage
                    WHERE payment_student_id = ?
                    ORDER BY period
                    """,
                    (student_id,),
                ).fetchall()
            )
            result.append(
                {
                    "id": int(student_id),
                    "name": name,
                    "monthly_amount": float(monthly_amount or 0),
                    "messages_enabled": bool(messages_enabled),
                    "coverage": coverage,
                }
            )
        return result


def _paid_through(student):
    paid_through = None
    for period, month_name in PAYMENT_PERIODS:
        if period not in student["coverage"]:
            break
        paid_through = (period, month_name)
    return paid_through


def _payment_status():
    ensure_payment_tables()
    with sqlite3.connect(bot.COREAPP_DB_PATH) as conn:
        imported = conn.execute(
            """
            SELECT imported_at, source_filename, student_count, coverage_count
            FROM payment_import_status WHERE id = 1
            """
        ).fetchone()
    return imported


def payments_summary_text():
    imported_at, _filename, _student_count, _coverage_count = _payment_status()
    students = _payment_rows()
    if not imported_at or not students:
        return (
            "💳 <b>Оплаты 11 класса</b>\n\n"
            "Данные ещё не загружены. Пришли мне в личный чат файл Excel с вкладкой «11 КЛАСС».\n\n"
            "Я прочитаю только эту вкладку. Другие листы и исходный файл не изменяются."
        )

    current = _current_course_period()
    period_name = dict(PAYMENT_PERIODS).get(current, "Вне учебного периода")
    current_paid = [student for student in students if current and current in student["coverage"]]
    expected = sum(student["monthly_amount"] for student in students)
    current_amount = sum(student["coverage"].get(current, 0) for student in students) if current else 0
    fully_paid = [
        student
        for student in students
        if _paid_through(student) and _paid_through(student)[0] == PAYMENT_PERIODS[-1][0]
    ]

    past_periods = _completed_course_periods()
    debt_entries = [
        (student, period)
        for student in students
        for period in past_periods
        if period not in student["coverage"]
    ]
    debt_amount = sum(student["monthly_amount"] for student, _period in debt_entries)

    try:
        shown_time = datetime.fromisoformat(imported_at).strftime("%d.%m.%Y %H:%M")
    except Exception:
        shown_time = str(imported_at)[:16]

    lines = [
        "💳 <b>Оплаты 11 класса</b>",
        "",
        f"Данные обновлены: {shown_time}",
        f"Учеников: <b>{len(students)}</b>",
    ]
    if current:
        lines.extend(
            [
                f"Текущий месяц: <b>{period_name.lower()}</b>",
                f"✅ Есть отметка об оплате: <b>{len(current_paid)} из {len(students)}</b>",
                f"Получено за месяц: <b>{_money(current_amount)}</b>",
                f"План по индивидуальным суммам: <b>{_money(expected)}</b>",
            ]
        )
    lines.extend(
        [
            f"🔴 Пропуски за завершённые месяцы: <b>{len(debt_entries)}</b> на {_money(debt_amount)}",
            f"🏁 Оплачено по май: <b>{len(fully_paid)}</b>",
            "",
            "🔒 Тестовый режим: сообщения ученикам не отправляются.",
        ]
    )
    return "\n".join(lines)


def payments_markup():
    students = _payment_rows()
    rows = []
    if students:
        rows.append([InlineKeyboardButton("👥 По ученикам", callback_data="cab:payments:list")])
        rows.append([InlineKeyboardButton("🔄 Обновить сводку", callback_data="cab:payments")])
    rows.append([InlineKeyboardButton("📎 Как обновить данные", callback_data="cab:payments:help")])
    rows.append([InlineKeyboardButton("← В кабинет", callback_data="cab:back")])
    return InlineKeyboardMarkup(rows)


def payment_students_markup():
    rows = []
    for student in _payment_rows():
        paid_through = _paid_through(student)
        if paid_through and paid_through[0] == PAYMENT_PERIODS[-1][0]:
            icon = "🏁"
        elif paid_through:
            icon = "✅"
        else:
            icon = "⚪"
        through_text = paid_through[1].lower() if paid_through else "нет отметок"
        label = f"{icon} {student['name']} · до {through_text}"
        if len(label) > 54:
            label = label[:53] + "…"
        rows.append(
            [InlineKeyboardButton(label, callback_data=f"cab:payments:student:{student['id']}")]
        )
    rows.append([InlineKeyboardButton("← К сводке", callback_data="cab:payments")])
    return InlineKeyboardMarkup(rows)


def payment_student_text(student_id):
    student = next((row for row in _payment_rows() if row["id"] == int(student_id)), None)
    if not student:
        return "Ученик не найден. Обнови файл с оплатами."
    lines = [
        f"💳 <b>{html.escape(student['name'])}</b>",
        f"Сумма за месяц: <b>{_money(student['monthly_amount'])}</b>",
        "",
    ]
    for period, month_name in PAYMENT_PERIODS:
        amount = student["coverage"].get(period)
        lines.append(f"✅ {month_name}: {_money(amount)}" if amount else f"▫️ {month_name}: нет отметки")
    paid_through = _paid_through(student)
    if paid_through and paid_through[0] == PAYMENT_PERIODS[-1][0]:
        lines.extend(["", "🏁 Курс оплачен по май. Напоминания не требуются."])
    lines.extend(["", "🔒 Сообщения ученику отключены на время теста."])
    return "\n".join(lines)


_previous_cabinet_markup = live23.cabinet_markup


def cabinet_markup_with_payments():
    base = _previous_cabinet_markup()
    rows = [list(row) for row in base.inline_keyboard]
    if not any(
        getattr(button, "callback_data", None) == "cab:payments"
        for row in rows
        for button in row
    ):
        rows.insert(max(0, len(rows) - 1), [
            InlineKeyboardButton("💳 Оплаты", callback_data="cab:payments")
        ])
    return InlineKeyboardMarkup(rows)


live23.cabinet_markup = cabinet_markup_with_payments

_previous_cabinet_callback = live23.cabinet_callback


async def cabinet_callback_with_payments(update, context):
    query = update.callback_query
    if not query or not live23._admin_private(update):
        return
    data = str(query.data or "")

    if data == "cab:payments":
        await query.answer()
        await query.edit_message_text(
            payments_summary_text(), parse_mode="HTML", reply_markup=payments_markup()
        )
        return

    if data == "cab:payments:list":
        await query.answer()
        await query.edit_message_text(
            "👥 <b>Оплаты по ученикам</b>\n\n"
            "Период «до …» показывает последний подряд оплаченный месяц.",
            parse_mode="HTML",
            reply_markup=payment_students_markup(),
        )
        return

    if data.startswith("cab:payments:student:"):
        try:
            student_id = int(data.rsplit(":", 1)[1])
        except Exception:
            await query.answer("Не получилось определить ученика")
            return
        await query.answer()
        await query.edit_message_text(
            payment_student_text(student_id),
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("← К ученикам", callback_data="cab:payments:list")],
                [InlineKeyboardButton("← В кабинет", callback_data="cab:back")],
            ]),
        )
        return

    if data == "cab:payments:help":
        await query.answer()
        await query.edit_message_text(
            "📎 <b>Обновление оплат</b>\n\n"
            "1. Выгрузи актуальную таблицу в формате Excel (.xlsx).\n"
            "2. Пришли файл мне в личный чат.\n"
            "3. Я прочитаю только вкладку «11 КЛАСС» и обновлю сводку.\n\n"
            "Пустые будущие месяцы не считаются долгом. Пропуск становится долгом только после завершения месяца. "
            "Сообщения ученикам во время теста не отправляются.",
            parse_mode="HTML",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("← К сводке", callback_data="cab:payments")]
            ]),
        )
        return

    await _previous_cabinet_callback(update, context)


live23.cabinet_callback = cabinet_callback_with_payments

_previous_import_students_document = bot.import_students_document


async def import_payments_or_students_document(update, context):
    if (
        update.effective_chat.type != "private"
        or not bot.user_is_admin(update)
        or not update.message
        or not update.message.document
    ):
        await _previous_import_students_document(update, context)
        return
    document = update.message.document
    if not str(document.file_name or "").lower().endswith(".xlsx"):
        await _previous_import_students_document(update, context)
        return

    telegram_file = await document.get_file()
    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)
    try:
        students = parse_payment_workbook(buffer.getvalue())
    except Exception as exc:
        await update.message.reply_text(f"Не получилось прочитать оплаты: {exc}")
        return
    if students is None:
        await _previous_import_students_document(update, context)
        return

    student_count, coverage_count = save_payment_snapshot(students, document.file_name)
    await update.message.reply_text(
        "✅ <b>Оплаты 11 класса обновлены</b>\n\n"
        f"Учеников: {student_count}\n"
        f"Заполненных месяцев: {coverage_count}\n\n"
        "Открой «👩‍🏫 Кабинет Маши» → «💳 Оплаты».\n"
        "🔒 Ученикам ничего не отправлено.",
        parse_mode="HTML",
    )


bot.import_students_document = import_payments_or_students_document


if __name__ == "__main__":
    live79.live71.ensure_molar_access_tables()
    live79.live70.ensure_health_tables()
    live79.live59.ensure_coreapp_webhook_audit_table()
    live79.live31.live25.ensure_lesson_day_before_table()
    live79.live31.live24.ensure_probnik_poll_tables()
    live79.live28.ensure_unanswered_reminder_tables()
    live79.live31.live30.live3.ensure_attendance_tables()
    live79.live31.live30.ensure_auto_attendance_table()
    live79.live31.ensure_personal_homework_reminder_table()
    live79.live17.ensure_acid_tables()
    live79.live24.live18.ensure_acid_reminder_table()
    live79.live34.ensure_attention_tables()
    live79.live35.ensure_admin_tasks_table()
    live79.ensure_task_sections()
    live79.live73.ensure_admin_task_view_state()
    live79.live35.seed_monday_task()
    live79.live37.update_monday_task_text()
    live79.live39.seed_probnik_return_task()
    live79.live41.ensure_weekly_report_tables()
    live79.live41.seed_current_trainers()
    live79.live48.ensure_metals_tables()
    live79.live48.register_metals_trainer()
    live79.live60.ensure_oxides_tables()
    live79.live60.register_oxides_trainer()
    live79.live43.ensure_probnik_analysis_tables()
    live79.live44.enable_probnik_analysis_now()
    live79.live46.ensure_monthly_auto_report_table()
    live79.live50.seed_molar_mass_task()
    live79.live51.ensure_course_schedule_table()
    live79.live66.seed_zlata_accounting_task()
    live79.live67.ensure_individual_students_table()
    live79.live71.complete_molar_mass_task()
    live79.live74.ensure_notification_catchup_tables()
    live79.live77.ensure_lesson_feedback_tables()
    live79.live78.seed_priority_tasks()
    live84.seed_teacher_product_tasks()
    ensure_payment_tables()
    live79.live56.log_probnik_cabinet_audit()
    print("Payment tracking ready: grade 11 Excel import, admin preview only", flush=True)
    print("Teacher product first five tasks seeded", flush=True)
    print("Admin tasks separated: active / completed / content / technical", flush=True)
    print("Admin current-task table shows first seven tasks", flush=True)
    print("Completed admin tasks stay completed after restart", flush=True)
    print("Lesson feedback ready: Mon/Wed 20:30, Sun 13:00; summary +1.5h", flush=True)
    print("Group traffic light ready", flush=True)
    print("Restart-safe notification catch-up enabled", flush=True)
    print("Admin task views auto-refresh after completion", flush=True)
    print("Today dashboard ready", flush=True)
    print("Molar mass calculator ready for admin and tutor", flush=True)
    print("Health monitoring and Telegram admin alerts enabled", flush=True)
    print("Probnik group reminders enabled: Thu/Fri + Friday poll + Sat morning", flush=True)
    print("Probnik personal no-response DMs enabled: 1.5h before probnik", flush=True)
    print("Probnik attention/parent escalation remains paused", flush=True)
    print("Individual students trainer-only mode ready", flush=True)
    print("Schedule-based homework cabinet ready", flush=True)
    print(f"Oxides trainer ready: questions={len(live79.live60.OXIDES_BANK)} monday_reminder=11:00", flush=True)
    print("CoreApp live sync receiver v2 ready", flush=True)
    live79.live24.main()
