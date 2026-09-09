import io
import json
import os
import sqlite3
import threading
from datetime import datetime, date, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

from phrases import DAILY_PHRASES

EXAM_DATE = date(2027, 6, 1)
PHRASE_START_DATE = date(2026, 9, 10)
TIMEZONE = ZoneInfo("Europe/Moscow")
COREAPP_DB_PATH = os.getenv("COREAPP_DB_PATH", "/tmp/coreapp_webhooks.db")

# Годовой курс: 108 уроков по понедельникам, средам и субботам.
# Первый урок — 07.09.2026. На 09.09.2026 по плану идёт урок №2.
COURSE_START_DATE = date(2026, 9, 7)
TOTAL_LESSONS = 108
LESSON_WEEKDAYS = {0, 2, 5}  # понедельник, среда, суббота
PROGRESS_SEGMENTS = 12


def today_moscow():
    return datetime.now(TIMEZONE).date()


def days_left():
    return (EXAM_DATE - today_moscow()).days


def get_daily_phrase(for_date=None):
    current_date = for_date or today_moscow()
    index = (current_date - PHRASE_START_DATE).days
    if 0 <= index < len(DAILY_PHRASES):
        return DAILY_PHRASES[index]
    return "Каждый день — ещё один маленький шаг к сотке."


def get_course_lesson_number(for_date=None):
    current_date = for_date or today_moscow()
    if current_date < COURSE_START_DATE:
        return 0

    count = 0
    cursor = COURSE_START_DATE
    while cursor <= current_date and count < TOTAL_LESSONS:
        if cursor.weekday() in LESSON_WEEKDAYS:
            count += 1
        cursor += timedelta(days=1)
    return min(count, TOTAL_LESSONS)


def get_course_progress_text(for_date=None):
    current_date = for_date or today_moscow()
    lesson_number = get_course_lesson_number(current_date)
    progress = lesson_number / TOTAL_LESSONS if TOTAL_LESSONS else 0
    percent = progress * 100

    if lesson_number == 0:
        filled = 0
    elif lesson_number >= TOTAL_LESSONS:
        filled = PROGRESS_SEGMENTS
    else:
        # На старте оставляем хотя бы одно розовое сердечко,
        # а точное значение всегда показываем цифрами рядом.
        filled = max(1, round(progress * PROGRESS_SEGMENTS))

    empty = PROGRESS_SEGMENTS - filled
    bar = "🩷" * filled + "🤍" * empty
    percent_text = f"{percent:.1f}".replace(".", ",")

    today_is_lesson = (
        current_date >= COURSE_START_DATE
        and current_date.weekday() in LESSON_WEEKDAYS
        and lesson_number <= TOTAL_LESSONS
    )

    lines = [
        "💗 <b>Прогресс курса</b>",
        bar,
        f"<b>{lesson_number} / {TOTAL_LESSONS} уроков</b> · {percent_text}%",
    ]

    if today_is_lesson and 0 < lesson_number <= TOTAL_LESSONS:
        lines.append(f"Сегодня — урок №{lesson_number} 🧪")

    return "\n".join(lines)


def get_countdown_text():
    days = days_left()
    progress_text = get_course_progress_text()

    if days > 1:
        return (
            "🧪 <b>ЕГЭ близко</b>\n\n"
            "До ЕГЭ по химии осталось\n"
            f"<b>{days} дней</b> 💗\n\n"
            f"{progress_text}\n\n"
            f"{get_daily_phrase()}"
        )

    if days == 1:
        return (
            "🧪 <b>ЕГЭ близко</b>\n\n"
            "До ЕГЭ по химии остался\n"
            "<b>1 день</b> 💗\n\n"
            f"{progress_text}\n\n"
            f"{get_daily_phrase()}"
        )

    if days == 0:
        return (
            "🧪 <b>ЕГЭ близко</b>\n\n"
            "<b>ЕГЭ ПО ХИМИИ — СЕГОДНЯ!</b> 💗\n\n"
            f"{progress_text}\n\n"
            "Вы уже сделали огромную работу. Теперь спокойно показываем всё, что умеем."
        )

    return (
        "🧪 <b>ЕГЭ близко</b>\n\n"
        f"{progress_text}\n\n"
        "ЕГЭ по химии уже позади 💗"
    )


def get_homework_reminder_text():
    return (
        "📝 <b>Время домашки</b> 💗\n\n"
        "Не откладываем на потом — проверьте, что сегодняшняя домашняя работа сделана и отправлена в CoreApp.\n\n"
        "Спокойно, системно и по чуть-чуть каждый день — так и приходят к сильному результату 🧪"
    )


def get_target_thread_id():
    thread_id = os.getenv("MESSAGE_THREAD_ID")
    return int(thread_id) if thread_id else None


def init_coreapp_db():
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS homework_submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                received_at TEXT NOT NULL,
                user_id TEXT,
                user_email TEXT,
                user_name TEXT,
                course_id TEXT,
                lesson_id TEXT,
                lesson_name TEXT,
                correct_count TEXT,
                total_count TEXT,
                raw_json TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                coreapp_user_id TEXT,
                user_email TEXT,
                user_name TEXT,
                course_id TEXT,
                active INTEGER NOT NULL DEFAULT 1,
                telegram_user_id INTEGER,
                telegram_username TEXT,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.commit()


def normalize_email(value):
    return str(value or "").strip().lower()


def upsert_student(payload, active=1):
    coreapp_user_id = str(payload.get("user_id", "")).strip()
    email = normalize_email(payload.get("user_email"))
    name = str(payload.get("user_name", "")).strip()
    course_id = str(payload.get("course_id", "")).strip()
    now = datetime.now(TIMEZONE).isoformat()

    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        row = None
        if coreapp_user_id:
            row = conn.execute(
                "SELECT id FROM students WHERE coreapp_user_id = ? LIMIT 1",
                (coreapp_user_id,),
            ).fetchone()
        if row is None and email:
            row = conn.execute(
                "SELECT id FROM students WHERE lower(user_email) = ? LIMIT 1",
                (email,),
            ).fetchone()

        if row:
            conn.execute(
                """
                UPDATE students
                SET coreapp_user_id = CASE WHEN ? != '' THEN ? ELSE coreapp_user_id END,
                    user_email = CASE WHEN ? != '' THEN ? ELSE user_email END,
                    user_name = CASE WHEN ? != '' THEN ? ELSE user_name END,
                    course_id = CASE WHEN ? != '' THEN ? ELSE course_id END,
                    active = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    coreapp_user_id, coreapp_user_id,
                    email, email,
                    name, name,
                    course_id, course_id,
                    active,
                    now,
                    row[0],
                ),
            )
        else:
            conn.execute(
                """
                INSERT INTO students (
                    coreapp_user_id, user_email, user_name, course_id,
                    active, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                (coreapp_user_id, email, name, course_id, active, now),
            )
        conn.commit()


def save_coreapp_submission(payload):
    upsert_student(payload, active=1)
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        conn.execute(
            """
            INSERT INTO homework_submissions (
                received_at, user_id, user_email, user_name, course_id,
                lesson_id, lesson_name, correct_count, total_count, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now(TIMEZONE).isoformat(),
                str(payload.get("user_id", "")),
                normalize_email(payload.get("user_email")),
                str(payload.get("user_name", "")),
                str(payload.get("course_id", "")),
                str(payload.get("lesson_id", "")),
                str(payload.get("lesson_name", "")),
                str(payload.get("correct_count", "")),
                str(payload.get("total_count", "")),
                json.dumps(payload, ensure_ascii=False),
            ),
        )
        conn.commit()


def get_recent_coreapp_submissions(limit=5):
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        return conn.execute(
            """
            SELECT received_at, user_name, user_email, lesson_name
            FROM homework_submissions
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()


def get_student_counts():
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        total = conn.execute("SELECT COUNT(*) FROM students WHERE active = 1").fetchone()[0]
        linked = conn.execute(
            "SELECT COUNT(*) FROM students WHERE active = 1 AND telegram_user_id IS NOT NULL"
        ).fetchone()[0]
    return total, linked


def get_latest_lesson():
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        row = conn.execute(
            """
            SELECT lesson_id, lesson_name
            FROM homework_submissions
            WHERE coalesce(lesson_id, '') != ''
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()
    return row


def get_homework_status_for_lesson(lesson_id):
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        students = conn.execute(
            """
            SELECT coreapp_user_id, user_email, user_name
            FROM students
            WHERE active = 1
            ORDER BY lower(coalesce(user_name, user_email, ''))
            """
        ).fetchall()
        submissions = conn.execute(
            """
            SELECT user_id, lower(user_email)
            FROM homework_submissions
            WHERE lesson_id = ?
            """,
            (lesson_id,),
        ).fetchall()

    submitted_ids = {str(row[0] or "").strip() for row in submissions if row[0]}
    submitted_emails = {normalize_email(row[1]) for row in submissions if row[1]}
    done, missing = [], []
    for coreapp_user_id, email, name in students:
        label = (name or email or coreapp_user_id or "ученик").strip()
        is_done = (
            (coreapp_user_id and str(coreapp_user_id).strip() in submitted_ids)
            or (email and normalize_email(email) in submitted_emails)
        )
        (done if is_done else missing).append(label)
    return done, missing


def link_telegram_user(email, telegram_user_id, telegram_username, telegram_name):
    email = normalize_email(email)
    now = datetime.now(TIMEZONE).isoformat()
    with sqlite3.connect(COREAPP_DB_PATH) as conn:
        row = conn.execute(
            "SELECT id FROM students WHERE lower(user_email) = ? LIMIT 1",
            (email,),
        ).fetchone()
        if row:
            conn.execute(
                """
                UPDATE students
                SET telegram_user_id = ?, telegram_username = ?,
                    user_name = CASE WHEN coalesce(user_name, '') = '' THEN ? ELSE user_name END,
                    updated_at = ?
                WHERE id = ?
                """,
                (telegram_user_id, telegram_username, telegram_name, now, row[0]),
            )
        else:
            conn.execute(
                """
                INSERT INTO students (
                    user_email, user_name, active, telegram_user_id,
                    telegram_username, updated_at
                ) VALUES (?, ?, 1, ?, ?, ?)
                """,
                (email, telegram_name, telegram_user_id, telegram_username, now),
            )
        conn.commit()


def import_students_from_xlsx(file_bytes):
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return 0, "Файл пустой"

    header_index = None
    email_col = None
    name_col = None

    for i, row in enumerate(rows[:20]):
        normalized = [str(v or "").strip().lower() for v in row]
        for j, value in enumerate(normalized):
            if email_col is None and ("email" in value or "e-mail" in value or "почт" in value):
                email_col = j
                header_index = i
            if name_col is None and ("фио" in value or value == "имя" or "name" in value):
                name_col = j
        if email_col is not None:
            break

    if email_col is None:
        return 0, "Не нашла колонку с e-mail"

    imported = 0
    for row in rows[(header_index or 0) + 1:]:
        if email_col >= len(row):
            continue
        email = normalize_email(row[email_col])
        if not email or "@" not in email:
            continue
        name = ""
        if name_col is not None and name_col < len(row):
            name = str(row[name_col] or "").strip()
        upsert_student({"user_email": email, "user_name": name}, active=1)
        imported += 1
    return imported, None


def get_admin_id():
    value = os.getenv("ADMIN_TELEGRAM_ID", "").strip()
    return int(value) if value.isdigit() else None


def user_is_admin(update: Update):
    admin_id = get_admin_id()
    return bool(admin_id and update.effective_user and update.effective_user.id == admin_id)


class CoreAppWebhookHandler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        print("CoreApp HTTP:", format % args)

    def _send_json(self, status_code, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status_code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/health":
            self._send_json(200, {"ok": True, "service": "ege-blizko-timer-bot"})
            return
        self._send_json(404, {"ok": False, "error": "not_found"})

    def do_POST(self):
        parsed = urlparse(self.path)
        allowed_paths = {
            "/coreapp/homework-submitted",
            "/coreapp/student-joined",
            "/coreapp/student-removed",
        }
        if parsed.path not in allowed_paths:
            self._send_json(404, {"ok": False, "error": "not_found"})
            return

        expected_secret = os.getenv("COREAPP_WEBHOOK_SECRET", "")
        supplied_secret = parse_qs(parsed.query).get("secret", [""])[0]
        if not expected_secret or supplied_secret != expected_secret:
            self._send_json(401, {"ok": False, "error": "unauthorized"})
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            raw_body = self.rfile.read(content_length).decode("utf-8")
            payload = json.loads(raw_body or "{}")
        except (ValueError, UnicodeDecodeError, json.JSONDecodeError):
            self._send_json(400, {"ok": False, "error": "invalid_json"})
            return

        try:
            if parsed.path == "/coreapp/homework-submitted":
                save_coreapp_submission(payload)
                print("CoreApp homework submitted:", payload.get("user_name"), payload.get("lesson_name"))
            elif parsed.path == "/coreapp/student-joined":
                upsert_student(payload, active=1)
                print("CoreApp student joined:", payload.get("user_name"), payload.get("user_email"))
            else:
                upsert_student(payload, active=0)
                print("CoreApp student removed:", payload.get("user_name"), payload.get("user_email"))
        except Exception as exc:
            print("CoreApp save error:", repr(exc))
            self._send_json(500, {"ok": False, "error": "storage_error"})
            return

        self._send_json(200, {"ok": True})


def start_http_server():
    init_coreapp_db()
    port = int(os.getenv("PORT", "8080"))
    server = ThreadingHTTPServer(("0.0.0.0", port), CoreAppWebhookHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    print(f"HTTP-сервер CoreApp запущен на порту {port}")
    return server


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Привет! Я таймер курса «ЕГЭ близко» 🧪\n\n"
        "Мои команды:\n"
        "/ege — сколько дней до ЕГЭ\n"
        "/weeks — сколько недель до ЕГЭ\n"
        "/progress — прогресс годового курса\n"
        "/chatid — показать ID этого чата\n"
        "/threadid — показать ID текущей темы\n"
        "/myid — показать ваш Telegram ID\n"
        "/link email — привязать себя к CoreApp\n"
        "/corestatus — статус интеграции CoreApp\n"
        "/homeworkstatus — кто сдал последнее ДЗ\n"
        "/test — отправить тестовый отсчёт в группу курса"
    )


async def ege(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(get_countdown_text(), parse_mode="HTML")


async def weeks(update: Update, context: ContextTypes.DEFAULT_TYPE):
    days = days_left()
    if days <= 0:
        await update.message.reply_text("ЕГЭ уже наступил 🧪")
        return
    weeks_count = days // 7
    remainder = days % 7
    await update.message.reply_text(
        "🧪 <b>ЕГЭ близко</b>\n\n"
        "До ЕГЭ осталось примерно\n"
        f"<b>{weeks_count} недель и {remainder} дней</b> 💗",
        parse_mode="HTML",
    )


async def progress(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(get_course_progress_text(), parse_mode="HTML")


async def chatid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"ID этого чата:\n<code>{update.effective_chat.id}</code>",
        parse_mode="HTML",
    )


async def threadid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    current_thread_id = update.effective_message.message_thread_id
    if current_thread_id is None:
        await update.message.reply_text("Отправьте /threadid внутри нужной темы форума.")
        return
    await update.message.reply_text(
        f"ID этой темы:\n<code>{current_thread_id}</code>",
        parse_mode="HTML",
    )


async def myid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"Ваш Telegram ID:\n<code>{update.effective_user.id}</code>",
        parse_mode="HTML",
    )


async def link(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_chat.type != "private":
        await update.message.reply_text("Привязку лучше делать в личном чате со мной 💗")
        return
    if not context.args:
        await update.message.reply_text("Напишите так: /link ваша_почта@example.com")
        return
    email = normalize_email(context.args[0])
    if "@" not in email:
        await update.message.reply_text("Похоже, это не e-mail. Попробуйте ещё раз.")
        return
    user = update.effective_user
    full_name = " ".join(x for x in [user.first_name, user.last_name] if x).strip()
    link_telegram_user(email, user.id, user.username or "", full_name)
    await update.message.reply_text("✅ Готово! Я связал ваш Telegram с CoreApp.")


async def corestatus(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not user_is_admin(update):
        if get_admin_id() is None:
            await update.message.reply_text("Сначала нужно назначить администратора. Отправьте мне /myid.")
        else:
            await update.message.reply_text("Эта команда доступна только преподавателю.")
        return

    rows = get_recent_coreapp_submissions(5)
    total, linked = get_student_counts()
    lines = [
        "✅ CoreApp подключён",
        f"Учеников в базе: {total}",
        f"Привязали Telegram: {linked}",
    ]
    if rows:
        lines.append("\nПоследние события:")
        for _, user_name, user_email, lesson_name in rows:
            who = user_name or user_email or "ученик"
            lesson = lesson_name or "домашняя работа"
            lines.append(f"• {who} — {lesson}")
    else:
        lines.append("\nСобытий о сдаче ДЗ пока не было.")
    await update.message.reply_text("\n".join(lines))


async def homeworkstatus(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not user_is_admin(update):
        await update.message.reply_text("Эта команда доступна только преподавателю.")
        return
    lesson = get_latest_lesson()
    if not lesson:
        await update.message.reply_text("Пока не получено ни одной сдачи ДЗ из CoreApp.")
        return
    lesson_id, lesson_name = lesson
    done, missing = get_homework_status_for_lesson(lesson_id)
    lines = [
        f"📝 {lesson_name or 'Последнее домашнее задание'}",
        f"Сдали: {len(done)}",
        f"Не сдали: {len(missing)}",
    ]
    if done:
        lines.append("\n✅ Сдали:\n" + "\n".join(f"• {name}" for name in done))
    if missing:
        lines.append("\n⏳ Пока не сдали:\n" + "\n".join(f"• {name}" for name in missing))
    await update.message.reply_text("\n".join(lines))


async def import_students_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_chat.type != "private":
        return
    if not user_is_admin(update):
        return
    document = update.message.document
    if not document or not document.file_name.lower().endswith(".xlsx"):
        return
    telegram_file = await document.get_file()
    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)
    imported, error = import_students_from_xlsx(buffer.getvalue())
    if error:
        await update.message.reply_text(f"Не получилось импортировать список: {error}")
        return
    await update.message.reply_text(f"✅ Импортировано учеников: {imported}")


async def test(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = os.getenv("CHAT_ID")
    if not chat_id:
        await update.message.reply_text("CHAT_ID пока не настроен.")
        return
    target_chat_id = int(chat_id)
    await context.bot.send_message(
        chat_id=target_chat_id,
        message_thread_id=get_target_thread_id(),
        text="🧪 <b>ТЕСТ ТАЙМЕРА</b>\n\n" + get_countdown_text(),
        parse_mode="HTML",
    )
    if update.effective_chat.id != target_chat_id:
        await update.message.reply_text("✅ Тестовое сообщение отправлено в группу курса.")


async def daily_countdown(context: ContextTypes.DEFAULT_TYPE):
    chat_id = os.getenv("CHAT_ID")
    if not chat_id:
        return
    await context.bot.send_message(
        chat_id=int(chat_id),
        message_thread_id=get_target_thread_id(),
        text=get_countdown_text(),
        parse_mode="HTML",
    )


async def daily_homework_reminder(context: ContextTypes.DEFAULT_TYPE):
    chat_id = os.getenv("CHAT_ID")
    if not chat_id:
        return
    await context.bot.send_message(
        chat_id=int(chat_id),
        message_thread_id=get_target_thread_id(),
        text=get_homework_reminder_text(),
        parse_mode="HTML",
    )


def main():
    token = os.getenv("BOT_TOKEN")
    if not token:
        raise RuntimeError("Переменная BOT_TOKEN не установлена")

    start_http_server()
    application = Application.builder().token(token).build()

    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("ege", ege))
    application.add_handler(CommandHandler("weeks", weeks))
    application.add_handler(CommandHandler("progress", progress))
    application.add_handler(CommandHandler("chatid", chatid))
    application.add_handler(CommandHandler("threadid", threadid))
    application.add_handler(CommandHandler("myid", myid))
    application.add_handler(CommandHandler("link", link))
    application.add_handler(CommandHandler("corestatus", corestatus))
    application.add_handler(CommandHandler("homeworkstatus", homeworkstatus))
    application.add_handler(CommandHandler("test", test))
    application.add_handler(MessageHandler(filters.Document.ALL, import_students_document))

    application.job_queue.run_daily(
        daily_countdown,
        time=datetime.strptime("09:00", "%H:%M").time().replace(tzinfo=TIMEZONE),
    )
    application.job_queue.run_daily(
        daily_homework_reminder,
        time=datetime.strptime("19:00", "%H:%M").time().replace(tzinfo=TIMEZONE),
        days=(0, 2, 6),
    )

    print("Бот запущен")
    application.run_polling()


if __name__ == "__main__":
    main()
